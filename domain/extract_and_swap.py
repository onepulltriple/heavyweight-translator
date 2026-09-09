if 'IMPORT LIBRARIES, VARIABLES, AND FILE PATHS':
    import os
    import docx
    from copy import deepcopy
    from docx import Document
    from math import ceil
    from xml.sax.saxutils import escape, unescape

    from app import input_parameters as IP
    from domain import conditions_checks as CC
    from domain import constants as CONST
    from domain import preservation_operations as PRSVOP
    from domain import tagging_operations as TAGOP
    from domain import xml_operations as XMLOP
    from file_io import csv_write_operations as CSVW
    from file_io import dict_operations as DO
    from file_io import file_operations as FO
    from support import auxiliary_operations as AUXOP
    from support import progress_indication_operations as PRIND


def extract_or_swap_text_in_docx(
        file_path_dictionary, 
        step, 
        temp_translation_dict = None,
    ):
    if temp_translation_dict is None:
        temp_translation_dict = {}

    doc = Document(file_path_dictionary["file_path_to_source_document"])
    
    current_op_count = 0
    count_of_relevant_paragraphs = PRIND.count_relevant_paragraphs(doc, step) 
    
    newest_print_progress_threshold = ceil(
        IP.percentage_increment_to_report/100*count_of_relevant_paragraphs
    )
 
    print_progress_increment = newest_print_progress_threshold

    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                current_op_count += process_paragraph_and_runs_within_it(
                    temp_translation_dict, 
                    paragraph, 
                    step,
                ) 
                newest_print_progress_threshold = PRIND.indicate_progress(
                    temp_translation_dict, 
                    step, 
                    newest_print_progress_threshold, 
                    print_progress_increment, 
                    count_of_relevant_paragraphs, 
                    current_op_count,
                )
            for table in part.tables:
                current_op_count = process_table_cells(
                    temp_translation_dict, 
                    table, 
                    step, 
                    newest_print_progress_threshold, 
                    print_progress_increment, 
                    count_of_relevant_paragraphs, 
                    current_op_count,
                    ) 

    for paragraph in doc.paragraphs:
        current_op_count += process_paragraph_and_runs_within_it(
            temp_translation_dict, 
            paragraph, 
            step,
            ) 
        newest_print_progress_threshold = PRIND.indicate_progress(
            temp_translation_dict, 
            step, 
            newest_print_progress_threshold, 
            print_progress_increment, 
            count_of_relevant_paragraphs, 
            current_op_count,
            )

    for table in doc.tables:
        current_op_count = process_table_cells(
            temp_translation_dict, 
            table, 
            step, 
            newest_print_progress_threshold, 
            print_progress_increment, 
            count_of_relevant_paragraphs, 
            current_op_count,
            ) 

    if step == CONST.EXTRACT:
        preprocessing_dict = {}

        if os.path.isfile(
            file_path_dictionary["preprocessing_dict_file_path"]
        ):
            preprocessing_dict = DO.read_json_dictionary(
                file_path_dictionary["preprocessing_dict_file_path"]
            )

        if IP.target_lang_cult not in preprocessing_dict:
            preprocessing_dict = DO.extend_json_dictionary(preprocessing_dict,{
                IP.target_lang_cult:{
                    "regex_to_locate_bad_translations":
                        "the_corrected_text_to_replace_bad_translations",
                    "For demonstration":"For example"
                    }
                }
            )

        DO.write_dict_to_json(
            preprocessing_dict, file_path_dictionary["preprocessing_dict_file_path"]
        )

        DO.write_dict_to_json(
            temp_translation_dict, file_path_dictionary["TEMP_translation_dict_file_path"]
        )
        if len(temp_translation_dict) > 0:
            CSVW.write_translation_dict_to_csv_simplified(
                temp_translation_dict, 
                file_path_dictionary["source_language_plain_texts_file_path"],
            )
            FO.save_to_text_file(
                file_path_dictionary["target_language_translations_file_path"], 
                [], 
                "\n",
            )

            print(
                f"\nThe text file containing the untranslated source text "
                f"has been written to: \n"
                f"{file_path_dictionary["source_language_plain_texts_file_path"]}\n"
            )
            print(
                f"An empty text file awaiting translated text in the target language "
                f"has been written to: \n"
                f"{file_path_dictionary["target_language_translations_file_path"]}\n"
            )

        print(f"There were {len(temp_translation_dict)} {step} operations.")
        
    if step == CONST.SWAP:
        print(f"There were {current_op_count} {step} operations.")
        print("Saving translated document...")
        doc.save(file_path_dictionary["output_document_path"])
        doc.save(file_path_dictionary["output_document_path_with_datetime"])
    

def consolidate_runs(paragraph): 
    text_consolidator = ""
    previous_run = None
    index_of_run = -1

    for current_run_or_hyperlink, next_run_or_hyperlink in (
        AUXOP.pairwise_circular(paragraph.iter_inner_content())
    ):
        index_of_run += 1

        if not current_run_or_hyperlink.text:
            previous_run = current_run_or_hyperlink
            continue

        if isinstance(current_run_or_hyperlink, docx.text.hyperlink.Hyperlink):
            current_hyperlink = current_run_or_hyperlink

            for current_run in current_hyperlink.runs:
                if current_run.text:
                    text_consolidator += current_run.text
                    current_run.text = text_consolidator                   
                    text_consolidator = ""

            previous_run = current_run_or_hyperlink
            continue

        elif isinstance(current_run_or_hyperlink, docx.text.run.Run):
            current_run = current_run_or_hyperlink

            if (
                CC.the_current_run_has_an_R_character(current_run)
                or CC.button_like_formatting_starts_and_ends_in_the_current_run(
                    current_run, text_consolidator
                )
                or CC.button_like_formatting_starts_and_ends_in_the_next_run(
                    next_run_or_hyperlink, text_consolidator
                )
                or CC.weird_symbol_bracketed_by_blank_char_starts_in_the_current_run(
                    current_run, next_run_or_hyperlink, text_consolidator
                ) 
                or CC.weird_symbol_bracketed_by_blank_char_ends_in_the_current_run(
                    previous_run, current_run, next_run_or_hyperlink, text_consolidator
                ) 
            ):
                text_consolidator += current_run.text
                current_run.text = text_consolidator                   
                text_consolidator = ""
                previous_run = current_run_or_hyperlink
                continue
            
            if (CC.either_has_special_characters(current_run, next_run_or_hyperlink)
                and (CC.the_current_run_has_one_or_two_special_characters(current_run)
                     or CC.the_next_run_has_one_or_two_special_characters(next_run_or_hyperlink))
                ):
                text_consolidator += current_run.text
                current_run.text = TAGOP.ignore_run_tag(index_of_run)
                
                if CC.the_last_run_in_the_paragraph_has_been_reached(next_run_or_hyperlink):
                    pass
                else:
                    previous_run = current_run_or_hyperlink
                    continue

            if (
                CC.button_like_formatting_starts_in_this_run(
                    current_run, next_run_or_hyperlink, text_consolidator
                )  
                or CC.button_like_formatting_ends_in_the_next_run(
                    next_run_or_hyperlink, text_consolidator
                )         
            ):
                text_consolidator += current_run.text
                current_run.text = TAGOP.ignore_run_tag(index_of_run)
                
                if CC.the_last_run_in_the_paragraph_has_been_reached(next_run_or_hyperlink):
                    pass
                else:
                    previous_run = current_run_or_hyperlink
                    continue

            if CC.bogus_change_of_nature_conditions_are_found(
                    previous_run, 
                    current_run, 
                    next_run_or_hyperlink,
                ):
                text_consolidator += current_run.text
                current_run.text = TAGOP.ignore_run_tag(index_of_run)
                
                if CC.the_last_run_in_the_paragraph_has_been_reached(next_run_or_hyperlink):
                    pass 
                else:
                    previous_run = current_run_or_hyperlink
                    continue

            if (
                CC.the_last_run_in_the_paragraph_has_been_reached(
                    next_run_or_hyperlink
                )
                or CC.there_is_no_text_in_the_next_run(
                    next_run_or_hyperlink
                )
                or CC.internal_hidden_text_style_has_been_reached(
                    next_run_or_hyperlink
                )
                or CC.button_like_formatting_starts_in_next_run(
                    next_run_or_hyperlink, text_consolidator
                )
                or CC.button_like_formatting_ends_in_this_run(
                    current_run, text_consolidator
                )
                or CC.there_is_a_change_of_nature(
                    current_run, next_run_or_hyperlink
                )
            ): 
                text_consolidator += current_run.text
                current_run.text = text_consolidator
                text_consolidator = ""                
                previous_run = current_run_or_hyperlink
                continue

            print(
                "The text-having run "
                f"\"{current_run_or_hyperlink.text}\" "
                "was not handled.\n"
            )

    return paragraph


def extract_runs(paragraph_with_cons_runs):
    index_of_run = -1
    paragraph_tagged_source_text_with_preserves = ""
    cons_run_tagged_text_with_preserves = ""

    previous_run = None

    for current_run_or_hyperlink in paragraph_with_cons_runs.iter_inner_content():
        index_of_run += 1

        if not current_run_or_hyperlink.text:
            current_glyph_holder = current_run_or_hyperlink
            current_glyph_holder.text = TAGOP.glyph_tag(index_of_run)
            cons_run_plain_text_with_preserves = current_glyph_holder.text
            cons_run_tagged_text_with_preserves = cons_run_plain_text_with_preserves
            cons_run_style = current_glyph_holder.style.name
            
        elif isinstance(current_run_or_hyperlink, docx.text.hyperlink.Hyperlink):
            current_hyperlink = current_run_or_hyperlink
            cons_run_escaped_text = escape(
                PRSVOP.pre_escape_preservations(
                    current_hyperlink.text
                )
            )
            cons_run_plain_text_with_preserves = (
                PRSVOP.preserve_run_special_items_with_temp_symbols(
                    cons_run_escaped_text
                )
            )
            cons_run_tagged_text_with_preserves = (
                TAGOP.hyperlink_tag(
                    cons_run_plain_text_with_preserves, 
                    index_of_run,
                )
            )
            cons_run_style = "Hyperlink"

        elif isinstance(current_run_or_hyperlink, docx.text.run.Run):
            current_run = current_run_or_hyperlink
            if (current_run.style.name != "Default Paragraph Font" 
                    and current_run.text != TAGOP.ignore_run_tag(index_of_run)):
                cons_run_escaped_text = escape(
                    PRSVOP.pre_escape_preservations(current_run.text)
                )
                cons_run_plain_text_with_preserves = (
                    PRSVOP.preserve_run_special_items_with_temp_symbols(
                        cons_run_escaped_text
                    )
                )
                cons_run_tagged_text_with_preserves = (
                    TAGOP.styled_run_tag(
                        cons_run_plain_text_with_preserves, 
                        index_of_run,
                    )
                )
            elif(CC.the_current_run_has_an_R_character(current_run)
                or CC.there_WAS_a_change_of_nature(current_run, previous_run)
                    and current_run.text != TAGOP.ignore_run_tag(index_of_run)
                    and not current_run.text.isspace()):
                cons_run_escaped_text = escape(
                    PRSVOP.pre_escape_preservations(current_run.text)
                )
                cons_run_plain_text_with_preserves = (
                    PRSVOP.preserve_run_special_items_with_temp_symbols(
                        cons_run_escaped_text
                    )
                )
                cons_run_tagged_text_with_preserves = TAGOP.changed_run_tag(
                    cons_run_plain_text_with_preserves, 
                    index_of_run,
                )
            elif(current_run.text == TAGOP.ignore_run_tag(index_of_run)):
                cons_run_plain_text_with_preserves = current_run.text
                cons_run_tagged_text_with_preserves = cons_run_plain_text_with_preserves
            else: 
                cons_run_escaped_text = (
                    escape(PRSVOP.pre_escape_preservations(current_run.text))
                )
                cons_run_plain_text_with_preserves = (
                    PRSVOP.preserve_run_special_items_with_temp_symbols(
                        cons_run_escaped_text
                    )
                )
                cons_run_tagged_text_with_preserves = cons_run_plain_text_with_preserves
            cons_run_style = current_run.style.name
        

        if (cons_run_style != "Default Paragraph Font"
            and cons_run_tagged_text_with_preserves != TAGOP.ignore_run_tag(index_of_run)):
            paragraph_tagged_source_text_with_preserves += cons_run_tagged_text_with_preserves
        else: 
            if cons_run_tagged_text_with_preserves == TAGOP.changed_run_tag(
                cons_run_plain_text_with_preserves,
                index_of_run,
            ):
                paragraph_tagged_source_text_with_preserves += cons_run_tagged_text_with_preserves
            elif cons_run_plain_text_with_preserves != TAGOP.ignore_run_tag(index_of_run):
                paragraph_tagged_source_text_with_preserves += cons_run_plain_text_with_preserves
            elif cons_run_plain_text_with_preserves == TAGOP.ignore_run_tag(index_of_run):
                paragraph_tagged_source_text_with_preserves += ""

        previous_run = current_run_or_hyperlink

    return paragraph_tagged_source_text_with_preserves


def paragraph_level_swapper(translation_dict, paragraph_with_cons_runs):
    from file_io.file_paths import file_path_dictionary
   
    carbon_copy_of_paragraph_with_cons_runs = deepcopy(paragraph_with_cons_runs)

    paragraph_tagged_source_text_with_preserves = (
        extract_runs(carbon_copy_of_paragraph_with_cons_runs)
    )
    
    if paragraph_tagged_source_text_with_preserves not in translation_dict:
        print(
            "The text element "
            f"\"{paragraph_tagged_source_text_with_preserves}\" "
            "was not found in the translation dictionary's keys."
        )
        return paragraph_with_cons_runs, 0
    
    if (paragraph_tagged_source_text_with_preserves != "" 
        and not paragraph_tagged_source_text_with_preserves.isspace()
        ):
        paragraph_tagged_translated_text_with_preserves = (
            translation_dict[paragraph_tagged_source_text_with_preserves][IP.target_lang_cult]
        )

    paragraph_tagged_translated_text = (
        PRSVOP.unpreserve_paragraph_translation(
            paragraph_tagged_translated_text_with_preserves
        )
    )
    translated_runs_with_tags = (
        XMLOP.split_string_into_list_of_tagged_and_untagged_elements(
            paragraph_tagged_translated_text
        )
    )

    if translated_runs_with_tags == file_path_dictionary["start_of_xml_debug_file_path"]:
        print(
            "Unparseable element encountered. "
            f"Review the element in \"{translated_runs_with_tags}\""
        )
        return paragraph_with_cons_runs, 0
    
    translated_paragraph = swap_runs(
        paragraph_with_cons_runs, 
        translated_runs_with_tags,
    ) 
    
    return translated_paragraph, 1


def swap_runs(paragraph_with_cons_runs, translated_runs_with_tags): 
    index_of_consolidated_run = -1
    index_of_translated_run = 0

    carbon_copy_of_paragraph_with_cons_runs = list(paragraph_with_cons_runs.iter_inner_content())
    carbon_copy_of_paragraph_with_cons_runs = deepcopy(carbon_copy_of_paragraph_with_cons_runs)

    for current_run_or_hyperlink in paragraph_with_cons_runs.iter_inner_content():
        index_of_consolidated_run += 1

        if index_of_translated_run < len(translated_runs_with_tags):      
            current_translated_run_dict = translated_runs_with_tags[index_of_translated_run]

            if ("type" in current_translated_run_dict.keys()
                and current_translated_run_dict["type"] == "glyph"):
                if (not current_run_or_hyperlink.text 
                    and index_of_consolidated_run == current_translated_run_dict["run_index"]
                    ):
                    current_glyph_holder = current_run_or_hyperlink
                    
                    index_of_translated_run += 1 
                else:
                    clear_cons_run_and_set_to_defaults(current_run_or_hyperlink)

            elif ("type" in current_translated_run_dict.keys()
                and current_translated_run_dict["type"] == "hyperlink"):
                if (isinstance(current_run_or_hyperlink, docx.text.hyperlink.Hyperlink)
                    and index_of_consolidated_run == current_translated_run_dict["run_index"]
                    ):
                    current_hyperlink = current_run_or_hyperlink                        
                    if (
                        "text" in current_translated_run_dict.keys()
                    ):
                        current_hyperlink.runs[0].text = unescape(
                            current_translated_run_dict["text"]
                        )
                    else:
                        print(
                            "Empty hyperlink encountered in: "
                            f"\"{paragraph_with_cons_runs.text}\"\n"
                            "Add text between the hyperlink tags."
                        )
                    index_of_translated_run += 1 
                else:
                    clear_cons_run_and_set_to_defaults(current_run_or_hyperlink)
                
            elif isinstance(current_run_or_hyperlink, docx.text.hyperlink.Hyperlink):
                print(
                    "Erroneous hyperlink encountered: "
                    f"\"{current_run_or_hyperlink.text}\" "
                    "Were the hyperlink tags dropped during translation?"
                )

            else:
                current_run = current_run_or_hyperlink
                if "text" in current_translated_run_dict.keys():
                    current_run.text = unescape(current_translated_run_dict["text"])
                else:
                    current_run.text = ""

                if "type" in current_translated_run_dict.keys():
                    index_of_styled_run = current_translated_run_dict["run_index"]

                    if (current_translated_run_dict["type"] == "styled"):
                        run_that_has_style_to_apply = (
                            carbon_copy_of_paragraph_with_cons_runs[index_of_styled_run]
                        )
                        current_run.style = run_that_has_style_to_apply.style
                    elif (current_translated_run_dict["type"] == "changed"):
                        run_that_has_changes_to_apply = (
                            carbon_copy_of_paragraph_with_cons_runs[index_of_styled_run]
                        )
                        current_run.font.color.rgb = run_that_has_changes_to_apply.font.color.rgb
                        current_run.font.size = run_that_has_changes_to_apply.font.size
                        current_run.font.name = run_that_has_changes_to_apply.font.name
                        current_run.style = run_that_has_changes_to_apply.style

                else:
                    current_run.style = "Default Paragraph Font" 
                    current_run.font.color.rgb = current_run._parent.style.font.color.rgb
                    current_run.font.size = None
                    current_run.font.name = None
                    current_run.font.subscript = None

                index_of_translated_run += 1 
        
        else: 
            clear_cons_run_and_set_to_defaults(current_run_or_hyperlink)

    return paragraph_with_cons_runs


def clear_cons_run_and_set_to_defaults(current_run_or_hyperlink):
    if not current_run_or_hyperlink.text:
        current_glyph_holder = current_run_or_hyperlink
        print(
            "Glyphs should not make it to this function. "
            f"See \"{current_run_or_hyperlink._parent.text}\".")
        return current_glyph_holder
        
    elif isinstance(current_run_or_hyperlink, docx.text.hyperlink.Hyperlink):
        current_hyperlink = current_run_or_hyperlink
        print(
            "Hyperlinks should not make it to this function. "
            f"See \"{current_run_or_hyperlink.text}\"."
        )
        return current_hyperlink

    elif isinstance(current_run_or_hyperlink, docx.text.run.Run):
        current_run = current_run_or_hyperlink        
        current_run.clear()
        current_run.style = "Default Paragraph Font" 
        return current_run


def process_table_cells(
    temp_translation_dict, 
    table, 
    step, 
    newest_print_progress_threshold, 
    print_progress_increment, 
    count_of_relevant_paragraphs, 
    current_op_count,
): 
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                current_op_count += process_paragraph_and_runs_within_it(
                    temp_translation_dict, 
                    paragraph, 
                    step,
                ) 
                newest_print_progress_threshold = PRIND.indicate_progress(
                    temp_translation_dict, 
                    step, 
                    newest_print_progress_threshold, 
                    print_progress_increment, 
                    count_of_relevant_paragraphs, 
                    current_op_count,
                )

            for nested_table in cell.tables:
                process_table_cells(
                    temp_translation_dict, 
                    nested_table, 
                    step, 
                    newest_print_progress_threshold, 
                    print_progress_increment, 
                    count_of_relevant_paragraphs, 
                    current_op_count,
                ) 

    return current_op_count


def process_paragraph_and_runs_within_it(temp_translation_dict, paragraph, step): 
    from file_io.dict_operations import maint_translation_dict
    
    if PRIND.is_relevant_paragraph(paragraph):
        cons_paragraph = consolidate_runs(paragraph) 

        if step == CONST.EXTRACT:
            paragraph_tagged_source_text_with_preserves = extract_runs(cons_paragraph)

            if (paragraph_tagged_source_text_with_preserves != "" 
                and not paragraph_tagged_source_text_with_preserves.isspace()
                and paragraph_tagged_source_text_with_preserves not in maint_translation_dict
                and paragraph_tagged_source_text_with_preserves not in temp_translation_dict
                ):
                temp_translation_dict[paragraph_tagged_source_text_with_preserves] = {
                    IP.target_lang_cult: None
                }
            
                return 0

        if step == CONST.SWAP:
            (paragraph, current_swap_count) = paragraph_level_swapper(
                maint_translation_dict, 
                cons_paragraph,
            )
            
            return current_swap_count
    
    if step == CONST.SWAP:
        return 0

    return -1